"""Excel (.xlsx) generator - openpyxl.

Sheet layout:
  Summary    - title, subtitle, metrics, source list
  Responses  - one row per question (number | question | answer), when in RFI mode
  <Tab name> - one sheet per remaining tab with its intro, bullets and table
"""

from __future__ import annotations

import io
import re
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .. import brand
from ..schemas import ResponseDocument, SourceInfo
from ._common import MEDIA_TYPES, output_filename

DARK_FILL = PatternFill("solid", fgColor=brand.office_hex(brand.DARK_BLUE))
GRAY_FILL = PatternFill("solid", fgColor=brand.office_hex(brand.GRAY))
HEADER_FONT = Font(name=brand.OFFICE_FONT, bold=True, size=11, color="FFFFFF")
TITLE_FONT = Font(name=brand.OFFICE_FONT, bold=True, size=18, color=brand.office_hex(brand.DARK_BLUE))
LABEL_FONT = Font(name=brand.OFFICE_FONT, bold=True, size=10, color=brand.office_hex(brand.DARK_BLUE))
BODY_FONT = Font(name=brand.OFFICE_FONT, size=10.5, color="33384F")
ACCENT_FONT = Font(name=brand.OFFICE_FONT, bold=True, size=14, color=brand.office_hex(brand.ORANGE))
VALUE_FONT = Font(name=brand.OFFICE_FONT, bold=True, size=14, color=brand.office_hex(brand.DARK_BLUE))
THIN = Side(style="thin", color="D8DEE8")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")

_INVALID_SHEET_CHARS = re.compile(r"[\[\]\*\?/\\:]")


def _sheet_title(name: str, used: set[str]) -> str:
    cleaned = _INVALID_SHEET_CHARS.sub("-", name).strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    counter = 2
    while candidate.lower() in used:
        suffix = f" ({counter})"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(candidate.lower())
    return candidate


def _widths(sheet: Worksheet, widths: Sequence[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _summary_sheet(sheet: Worksheet, doc: ResponseDocument, sources: Sequence[SourceInfo]) -> None:
    _widths(sheet, [34, 26, 26, 26])
    sheet["A1"] = doc.title
    sheet["A1"].font = TITLE_FONT
    row = 2
    if doc.subtitle:
        sheet["A2"] = doc.subtitle
        sheet["A2"].font = BODY_FONT
        sheet["A2"].alignment = WRAP_TOP
        row = 3

    row += 1
    sheet.cell(row=row, column=1, value="Mode").font = LABEL_FONT
    sheet.cell(row=row, column=2, value="RFI - questions answered" if doc.mode == "rfi" else "Summary dashboard").font = BODY_FONT
    row += 1
    if doc.mode == "rfi":
        sheet.cell(row=row, column=1, value="Questions answered").font = LABEL_FONT
        sheet.cell(row=row, column=2, value=doc.question_count).font = BODY_FONT
        row += 1

    if doc.metrics:
        row += 1
        sheet.cell(row=row, column=1, value="HEADLINE METRICS").font = LABEL_FONT
        row += 1
        for metric in doc.metrics:
            value_cell = sheet.cell(row=row, column=1, value=metric.value)
            value_cell.font = ACCENT_FONT if metric.accent else VALUE_FONT
            value_cell.fill = GRAY_FILL
            label_cell = sheet.cell(row=row, column=2, value=metric.label)
            label_cell.font = BODY_FONT
            label_cell.fill = GRAY_FILL
            row += 1

    if sources:
        row += 1
        sheet.cell(row=row, column=1, value="SOURCE DOCUMENTS").font = LABEL_FONT
        row += 1
        for header_index, header in enumerate(("File", "Characters", "Words", "Note"), start=1):
            cell = sheet.cell(row=row, column=header_index, value=header)
            cell.font = HEADER_FONT
            cell.fill = DARK_FILL
        row += 1
        for source in sources:
            sheet.cell(row=row, column=1, value=source.name).font = BODY_FONT
            sheet.cell(row=row, column=2, value=source.chars).font = BODY_FONT
            sheet.cell(row=row, column=3, value=source.words).font = BODY_FONT
            note = sheet.cell(row=row, column=4, value=source.note)
            note.font = BODY_FONT
            note.alignment = WRAP_TOP
            row += 1

    row += 1
    footer = sheet.cell(row=row, column=1, value=brand.footer_text())
    footer.font = Font(name=brand.OFFICE_FONT, size=8, color="8A90A8")


def _responses_sheet(sheet: Worksheet, doc: ResponseDocument) -> None:
    _widths(sheet, [10, 52, 96])
    for index, header in enumerate(("Ref", "Query", "Firstsource Response"), start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = HEADER_FONT
        cell.fill = DARK_FILL
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    row = 2
    for tab in doc.tabs:
        for item in tab.qa:
            for column, value in enumerate((item.n, item.q, item.a), start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.font = BODY_FONT
                cell.alignment = WRAP_TOP
                cell.border = CELL_BORDER
            sheet.row_dimensions[row].height = 58
            row += 1
    sheet.auto_filter.ref = f"A1:C{max(1, row - 1)}"


def _tab_sheet(sheet: Worksheet, tab) -> None:
    _widths(sheet, [40, 40, 40, 40, 40])
    row = 1
    sheet.cell(row=row, column=1, value=tab.name).font = TITLE_FONT
    row += 1
    if tab.intro:
        cell = sheet.cell(row=row, column=1, value=tab.intro)
        cell.font = BODY_FONT
        cell.alignment = WRAP_TOP
        row += 2

    if tab.bullets:
        sheet.cell(row=row, column=1, value="KEY POINTS").font = LABEL_FONT
        row += 1
        for bullet in tab.bullets:
            cell = sheet.cell(row=row, column=1, value=bullet)
            cell.font = BODY_FONT
            cell.alignment = WRAP_TOP
            row += 1
        row += 1

    if tab.table and tab.table.rows:
        if tab.table.headers:
            for index, header in enumerate(tab.table.headers, start=1):
                cell = sheet.cell(row=row, column=index, value=header)
                cell.font = HEADER_FONT
                cell.fill = DARK_FILL
            row += 1
        for data_row in tab.table.rows:
            for index, value in enumerate(data_row, start=1):
                cell = sheet.cell(row=row, column=index, value=value)
                cell.font = BODY_FONT
                cell.alignment = WRAP_TOP
                cell.border = CELL_BORDER
            row += 1
        row += 1

    if tab.callout and (tab.callout.title or tab.callout.body):
        title_cell = sheet.cell(row=row, column=1, value=tab.callout.title or "Key takeaway")
        title_cell.font = LABEL_FONT
        title_cell.fill = GRAY_FILL
        row += 1
        body_cell = sheet.cell(row=row, column=1, value=tab.callout.body)
        body_cell.font = BODY_FONT
        body_cell.alignment = WRAP_TOP
        body_cell.fill = GRAY_FILL


def generate(
    document_data: ResponseDocument,
    sources: Sequence[SourceInfo] | None = None,
) -> tuple[bytes, str, str]:
    sources = list(sources or [])
    workbook = Workbook()
    used: set[str] = set()

    summary = workbook.active
    summary.title = _sheet_title("Summary", used)
    _summary_sheet(summary, document_data, sources)

    if document_data.question_count:
        _responses_sheet(workbook.create_sheet(_sheet_title("Responses", used)), document_data)

    for tab in document_data.tabs:
        # Q&A already has a dedicated sheet; skip tabs that carry nothing else.
        if tab.qa and not (tab.bullets or tab.intro or tab.table or tab.callout):
            continue
        _tab_sheet(workbook.create_sheet(_sheet_title(tab.name, used)), tab)

    buffer = io.BytesIO()
    workbook.save(buffer)
    filename = output_filename(document_data.title, "xlsx")
    return buffer.getvalue(), filename, MEDIA_TYPES["xlsx"]
