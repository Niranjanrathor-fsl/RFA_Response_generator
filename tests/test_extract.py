"""Server-side extraction, replacing the browser extractor."""

from __future__ import annotations

import io

import pytest

from app.extract import ExtractedDocument, ExtractionError, build_corpus, extract_text


def test_plain_text():
    doc = extract_text("notes.txt", b"Q1. Describe your governance.\n\nQ2. And your labs?")
    assert "Q1." in doc.text
    assert doc.words > 5
    assert not doc.is_empty


def test_csv_becomes_pipe_delimited():
    doc = extract_text("data.csv", b"Question,Answer\nLiability?,Capped at 12 months\n")
    assert "Question | Answer" in doc.text


def test_json_is_pretty_printed():
    doc = extract_text("payload.json", b'{"q":"Do you build foundational models?"}')
    assert "foundational models" in doc.text


def test_docx_roundtrip():
    from docx import Document

    document = Document()
    document.add_paragraph("Q1. Describe your liability model.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Area"
    table.cell(0, 1).text = "Detail"
    buffer = io.BytesIO()
    document.save(buffer)

    doc = extract_text("rfi.docx", buffer.getvalue())
    assert "liability model" in doc.text
    assert "Area | Detail" in doc.text


def test_pptx_roundtrip_includes_slide_markers():
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "Kairos architecture overview"
    buffer = io.BytesIO()
    presentation.save(buffer)

    doc = extract_text("deck.pptx", buffer.getvalue())
    assert "Kairos architecture overview" in doc.text
    assert "--- Slide 1 ---" in doc.text


def test_xlsx_roundtrip_includes_sheet_markers():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Questions"
    sheet.append(["Ref", "Query"])
    sheet.append(["Q1", "Describe your labs"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    doc = extract_text("questions.xlsx", buffer.getvalue())
    assert "--- Sheet: Questions ---" in doc.text
    assert "Q1 | Describe your labs" in doc.text


def test_pdf_roundtrip():
    pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)

    doc = extract_text("scan.pdf", buffer.getvalue())
    # A blank page yields no text - the user must be told, not left guessing.
    assert doc.is_empty
    assert "scanned" in doc.note.lower()


def test_legacy_format_rejected_with_guidance():
    with pytest.raises(ExtractionError) as excinfo:
        extract_text("old.doc", b"\xd0\xcf\x11\xe0")
    assert "legacy" in str(excinfo.value).lower()


def test_unsupported_extension_rejected():
    with pytest.raises(ExtractionError):
        extract_text("archive.zip", b"PK\x03\x04")


def test_per_document_truncation(monkeypatch):
    from app.config import get_settings

    settings = get_settings()
    monkeypatch.setattr(settings, "max_chars_per_document", 100)
    doc = extract_text("long.txt", b"x" * 5000)
    assert len(doc.text) < 200
    assert "truncated" in doc.text


# --------------------------------------------------------------- corpus tests
def test_single_source_has_no_markers():
    corpus, truncated = build_corpus([ExtractedDocument("only.txt", "Just one document.")])
    assert corpus == "Just one document."
    assert "SOURCE DOCUMENT" not in corpus
    assert truncated is False


def test_multiple_sources_are_delimited_and_named():
    corpus, _ = build_corpus([
        ExtractedDocument("ClientRFI.docx", "Q1. Governance?"),
        ExtractedDocument("Capabilities.pptx", "Kairos overview."),
    ])
    assert "===== SOURCE DOCUMENT 1 of 2: ClientRFI.docx =====" in corpus
    assert "===== SOURCE DOCUMENT 2 of 2: Capabilities.pptx =====" in corpus


def test_empty_documents_are_dropped():
    corpus, _ = build_corpus([
        ExtractedDocument("good.txt", "Real content."),
        ExtractedDocument("blank.pdf", "   "),
    ])
    assert "blank.pdf" not in corpus
    assert "SOURCE DOCUMENT" not in corpus  # only one usable source remains


def test_budget_is_shared_evenly_so_one_file_cannot_crowd_out_another():
    corpus, truncated = build_corpus(
        [ExtractedDocument("big.txt", "a" * 5000), ExtractedDocument("small.txt", "b" * 5000)],
        max_total_chars=1000,
    )
    assert truncated is True
    assert "big.txt" in corpus and "small.txt" in corpus
    assert corpus.count("document truncated") == 2


def test_no_usable_documents_returns_empty():
    corpus, truncated = build_corpus([ExtractedDocument("blank.txt", "")])
    assert corpus == ""
    assert truncated is False


def test_html_extracts_visible_text_and_drops_script_and_style():
    html = (
        b"<html><head><style>.a{color:red}</style>"
        b"<script>var secret = 'do not index me';</script></head>"
        b"<body><h1>Kairos proposal studio</h1><p>Governed Autonomy applies.</p></body></html>"
    )
    doc = extract_text("proposal-studio.html", html)
    assert "Kairos proposal studio" in doc.text
    assert "Governed Autonomy applies." in doc.text
    assert "do not index me" not in doc.text
    assert "color:red" not in doc.text


def test_html_tables_become_pipe_delimited_rows():
    html = (
        b"<html><body><table>"
        b"<tr><th>Area</th><th>Detail</th></tr>"
        b"<tr><td>Generative AI</td><td>32-38 engagements</td></tr>"
        b"</table></body></html>"
    )
    doc = extract_text("capabilities.html", html)
    assert "Area | Detail" in doc.text
    assert "Generative AI | 32-38 engagements" in doc.text


def test_legacy_xls_is_now_readable():
    import xlwt  # test-only writer for the legacy format

    book = xlwt.Workbook()
    sheet = book.add_sheet("Pricing")
    sheet.write(0, 0, "Tier")
    sheet.write(0, 1, "Monthly Fee")
    sheet.write(1, 0, "Starter")
    sheet.write(1, 1, "15000")
    buffer = io.BytesIO()
    book.save(buffer)

    doc = extract_text("pricing.xls", buffer.getvalue())
    assert "Tier | Monthly Fee" in doc.text
    assert "Starter" in doc.text


def test_xls_named_file_that_is_really_xlsx_is_readable():
    # "Forrester responses.XLS" is a modern workbook under a legacy extension;
    # xlrd rejects it with "Excel xlsx file; not supported".
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Responses"
    sheet.append(["Question", "Answer"])
    sheet.append(["Headcount?", "About 30,000"])
    buffer = io.BytesIO()
    book.save(buffer)

    doc = extract_text("Forrester responses.XLS", buffer.getvalue())
    assert "Question | Answer" in doc.text
    assert "About 30,000" in doc.text


def test_rejected_binary_formats_yield_no_sections():
    from app.config import get_settings
    from app.document_sections import extract_sections

    assert extract_sections("clip.mp4", b"\x00\x00\x00\x18ftypmp42", get_settings()) == []
    assert extract_sections("bundle.zip", b"PK\x03\x04rubbish", get_settings()) == []


def test_unknown_extension_that_looks_like_text_is_read():
    from app.config import get_settings
    from app.document_sections import extract_sections

    sections = extract_sections("config.xml", b"<root><note>Kairos</note></root>", get_settings())
    assert sections
    assert "Kairos" in sections[0][1]


def test_unknown_extension_that_is_binary_is_skipped():
    from app.config import get_settings
    from app.document_sections import extract_sections

    assert extract_sections("thing.bin", b"\x00\x01\x02\x03\xff\xfe", get_settings()) == []


# ------------------------------------------------------------ image uploads
def _png_bytes(width=80, height=60):
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (10, 80, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


def test_image_formats_are_accepted_for_upload():
    from app.extract import SUPPORTED_EXTENSIONS

    for ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"):
        assert ext in SUPPORTED_EXTENSIONS, f"{ext} should be uploadable"


def test_image_upload_is_read_by_vision(monkeypatch):
    """A chart or screenshot must become text, the same way it already does on
    the RAG ingestion path. Vision is forced off suite-wide by conftest, so this
    test turns it on for itself and stubs the model call."""
    import app.document_sections as ds
    import app.extract as ex
    from app.config import get_settings

    vision_on = get_settings().model_copy(update={"rag_vision_enabled": True})
    monkeypatch.setattr(ex, "get_settings", lambda: vision_on)
    monkeypatch.setattr(ds, "describe_or_skip", lambda *a, **k: "Bar chart: FY23 120, FY24 210.")

    doc = extract_text("chart.png", _png_bytes())
    assert "FY24 210" in doc.text
    assert not doc.is_empty


def test_unreadable_image_explains_itself_rather_than_failing(monkeypatch):
    """Vision off, or a purely decorative image: the user needs to know why
    nothing came back, not a bare 'no readable text' error."""
    import app.document_sections as ds

    monkeypatch.setattr(ds, "describe_or_skip", lambda *a, **k: None)
    doc = extract_text("logo.png", _png_bytes())
    assert doc.is_empty
    assert "image" in doc.note.lower()


def _docx_with_repeated_picture(copies: int, distinct_png: bytes | None = None) -> bytes:
    """A Word file holding the same picture `copies` times under DIFFERENT part
    names - what form-style documents do for every checkbox. python-docx itself
    would share one part, so the copies are related by hand."""
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.parts.image import ImagePart

    document = Document()
    document.add_paragraph("Q1. Do you support 24x7 delivery?")
    blob = _png_bytes()
    for i in range(copies):
        part = ImagePart(PackURI(f"/word/media/checkbox{i}.png"), "image/png", blob, document.part.package)
        document.part.relate_to(part, RT.IMAGE)
    if distinct_png is not None:
        part = ImagePart(PackURI("/word/media/chart.png"), "image/png", distinct_png, document.part.package)
        document.part.relate_to(part, RT.IMAGE)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_a_picture_repeated_under_different_names_is_read_once(monkeypatch):
    import app.document_sections as ds
    from app.config import get_settings

    calls = []

    def fake_vision(blob, **kwargs):
        calls.append(blob)
        return f"Picture {len(calls)}"

    monkeypatch.setattr(ds, "describe_or_skip", fake_vision)
    settings = get_settings().model_copy(update={"rag_vision_enabled": True})

    data = _docx_with_repeated_picture(copies=50, distinct_png=_png_bytes(120, 90))
    sections = ds.extract_docx_sections(data, settings)

    assert len(calls) == 2  # the checkbox once, the different chart once
    assert [text for label, text in sections if label.startswith("Image")] == ["Picture 1", "Picture 2"]
    assert "24x7" in sections[0][1]


def test_config_groups_extensions_for_display():
    """The UI showed the alphabetically-first 8 extensions, which is meaningless
    to a user. The server should hand over a curated, ordered list instead."""
    from app.routes.health import client_config

    config = client_config()
    assert "accepted_display" in config
    display = config["accepted_display"]
    assert display[0].startswith("PDF")
    joined = " ".join(display)
    assert "Word" in joined and "Images" in joined
