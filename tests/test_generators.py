"""Every output format must build real, openable files with brand styling."""

from __future__ import annotations

import io
import zipfile

import pytest

from app.generators import FORMAT_LABELS, generate
from app.schemas import SourceInfo

SOURCES = [
    SourceInfo(name="ClientRFI.docx", chars=1200, words=180),
    SourceInfo(name="Capabilities.pptx", chars=800, words=110, note="Truncated for the model."),
]


@pytest.mark.parametrize("fmt", sorted(FORMAT_LABELS))
def test_every_format_produces_a_non_trivial_file(fmt, rfi_document):
    content, filename, media_type = generate(fmt, rfi_document, sources=SOURCES)
    assert len(content) > 1000, f"{fmt} output suspiciously small"
    assert filename.endswith(("html", "docx", "pptx", "xlsx"))
    assert "acme-bank" in filename  # slugified from the title
    assert media_type


@pytest.mark.parametrize("fmt", ["docx", "pptx", "xlsx"])
def test_office_files_are_valid_ooxml_packages(fmt, rfi_document):
    content, _, _ = generate(fmt, rfi_document, sources=SOURCES)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert archive.testzip() is None
        names = archive.namelist()
    assert "[Content_Types].xml" in names


def test_dashboard_html_is_self_contained_and_on_brand(rfi_document):
    content, filename, media_type = generate("dashboard", rfi_document, sources=SOURCES)
    html = content.decode("utf-8")
    assert media_type.startswith("text/html")
    assert filename.endswith(".html")
    # Brand palette and font stack
    assert "#1E2247" in html and "#DF6014" in html
    assert "Neue Haas Grotesk Display" in html
    # Real logo, embedded rather than linked, so the file travels
    assert "data:image/png;base64," in html
    assert "http://" not in html.replace("http://www.w3.org", "")  # no external fetches
    # Footer year and tagline
    assert "Firstsource. All rights reserved." in html
    # Every tab has a button and a matching panel
    for index in range(len(rfi_document.tabs)):
        assert f'data-target="tab-{index}"' in html
        assert f'id="tab-{index}"' in html


def test_dashboard_css_is_not_html_escaped(rfi_document):
    """Brand CSS lives inside <style>; HTML-escaping it breaks the font stack."""
    html = generate("dashboard", rfi_document)[0].decode("utf-8")
    style = html.split("<style>", 1)[1].split("</style>", 1)[0]
    assert "&#39;" not in style, "single quotes were HTML-escaped inside <style>"
    assert "&quot;" not in style
    assert "font-family:'Neue Haas Grotesk Display'" in style
    assert "format('truetype')" in style


def test_dashboard_renders_all_questions_with_original_labels(rfi_document):
    html = generate("dashboard", rfi_document, sources=SOURCES)[0].decode("utf-8")
    assert "Q1" in html and "Q7" in html and "3.2(a)" in html
    assert "capped at 12 months" in html.lower()


def test_qa_document_lists_questions_and_a_contents_block(rfi_document):
    html = generate("qa", rfi_document, sources=SOURCES)[0].decode("utf-8")
    assert "3 questions answered" in html
    assert "Contents" in html
    assert "foundational LLMs" in html


def test_summary_mode_has_no_question_furniture(summary_document):
    html = generate("dashboard", summary_document)[0].decode("utf-8")
    assert summary_document.mode == "summary"
    assert "Architecture" in html and "Outcomes" in html
    assert "questions answered" not in html


def test_html_escapes_injected_markup(summary_document):
    summary_document.title = '<script>alert("xss")</script>'
    html = generate("dashboard", summary_document)[0].decode("utf-8")
    assert "<script>alert" not in html
    assert "&lt;script&gt;" in html


def test_word_document_contains_the_answers(rfi_document):
    import docx

    content, _, _ = generate("docx", rfi_document, sources=SOURCES)
    document = docx.Document(io.BytesIO(content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert rfi_document.title in text
    assert "E&O and cyber-liability insurance" in text
    table_text = " ".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "Generative AI" in table_text


def test_powerpoint_has_a_slide_per_section_plus_title_and_metrics(rfi_document):
    from pptx import Presentation

    content, _, _ = generate("pptx", rfi_document, sources=SOURCES)
    presentation = Presentation(io.BytesIO(content))
    assert len(presentation.slides) >= len(rfi_document.tabs) + 2
    all_text = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_text.append(shape.text_frame.text)
    joined = "\n".join(all_text)
    assert rfi_document.title in joined
    assert "Innovation labs" in joined or "INNOVATION LABS" in joined
    assert "Q1" in joined


def test_excel_has_summary_and_responses_sheets(rfi_document):
    from openpyxl import load_workbook

    content, _, _ = generate("xlsx", rfi_document, sources=SOURCES)
    workbook = load_workbook(io.BytesIO(content))
    assert "Summary" in workbook.sheetnames
    assert "Responses" in workbook.sheetnames

    responses = workbook["Responses"]
    assert [c.value for c in responses[1]] == ["Ref", "Query", "Firstsource Response"]
    refs = [responses.cell(row=r, column=1).value for r in range(2, responses.max_row + 1)]
    assert refs == ["Q1", "Q7", "3.2(a)"]

    summary = workbook["Summary"]
    flat = " ".join(str(c.value) for row in summary.iter_rows() for c in row if c.value)
    assert "ClientRFI.docx" in flat
    assert "Firstsource. All rights reserved." in flat


def test_summary_mode_excel_omits_the_responses_sheet(summary_document):
    from openpyxl import load_workbook

    content, _, _ = generate("xlsx", summary_document)
    workbook = load_workbook(io.BytesIO(content))
    assert "Responses" not in workbook.sheetnames


def test_long_tab_names_are_truncated_for_excel(summary_document):
    from openpyxl import load_workbook

    summary_document.tabs[1].name = "A very long section name that exceeds Excel's limit of 31"
    content, _, _ = generate("xlsx", summary_document)
    workbook = load_workbook(io.BytesIO(content))
    assert all(len(name) <= 31 for name in workbook.sheetnames)


def test_unknown_format_raises(rfi_document):
    with pytest.raises(KeyError):
        generate("pdf", rfi_document)
